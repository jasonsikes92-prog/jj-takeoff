"""Read returned trade response workbooks without changing estimate prices."""
import argparse
import hashlib
import json
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from quote_pricing_review import number

HEADERS = ['Scope ID', 'Work', 'Reference quantity', 'Reference unit', 'Response',
           'Included under scope ID', 'Quoted quantity', 'Quoted unit',
           'Unit price (USD)', 'Line amount (USD)', 'Clarification / exclusions']
METADATA = {'company': 'B5', 'quote_number': 'D5', 'contact': 'B6', 'quote_date': 'D6',
            'lead_time': 'B7', 'work_duration': 'D7', 'delivery_included': 'B8',
            'tax_included': 'D8', 'delivery_charge': 'B9', 'tax_charge': 'D9',
            'quoted_total': 'B10', 'price_change_conditions': 'D10',
            'other_charges_exclusions': 'B11', 'additional_work_options': 'D11'}
STATUSES = {'Priced separately', 'Included in package', 'Excluded',
            'Needs clarification', 'Reference only'}


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def review_response(scope, metadata, responses, current_snapshot):
    if scope.get('snapshot_sha256') != current_snapshot:
        raise ValueError('Scope differs from the current estimate snapshot; renew the response review')
    expected = {i['row_id']: i for i in scope['items']}
    ids = [r['scope_id'] for r in responses]
    if len(expected) != len(scope['items']) or len(ids) != len(set(ids)) or set(ids) != set(expected):
        raise ValueError('Returned response must preserve every unique scope ID')
    findings = []
    reviewed = []
    charged = {}
    by_id = {r['scope_id']: r for r in responses}
    for field in ('company', 'quote_number', 'quote_date'):
        if blank(metadata.get(field)):
            findings.append('Missing quote field: ' + field)
    for response in responses:
        identity = response['scope_id']; item = expected[identity]; issues = []
        status = response.get('response')
        if blank(status):
            issues.append('No response supplied')
        elif status not in STATUSES:
            issues.append('Unknown response status')
        elif status in ('Excluded', 'Needs clarification'):
            issues.append('Scope requires resolution: ' + status)
        values = {}
        for key in ('quoted_quantity', 'unit_price', 'line_amount'):
            try:
                values[key] = None if blank(response.get(key)) else number(response[key])
            except ValueError:
                values[key] = None
                issues.append('Invalid nonnegative numeric field: ' + key)
        parent = response.get('included_under')
        if status == 'Included in package':
            if parent not in by_id or parent == identity or by_id[parent].get('response') != 'Priced separately':
                issues.append('Included scope needs a different, separately priced scope ID from this trade')
            if any(not blank(response.get(k)) for k in ('quoted_quantity', 'quoted_unit', 'unit_price', 'line_amount')):
                issues.append('Included scope has separate billing fields; possible duplicate charge')
        elif not blank(parent):
            issues.append('Package owner supplied without Included in package status')
        if status == 'Priced separately':
            quantity, rate, amount = (values[k] for k in ('quoted_quantity', 'unit_price', 'line_amount'))
            if amount is None:
                issues.append('Line amount is missing or invalid')
            if quantity is None or rate is None or blank(response.get('quoted_unit')):
                issues.append('Billing quantity, unit and rate need clarification; no lump sum inferred')
            elif amount is not None and (quantity * rate).quantize(Decimal('.01'), rounding=ROUND_HALF_UP) != amount:
                issues.append('Quoted quantity times rate differs from line amount')
            if not blank(response.get('quoted_unit')) and response['quoted_unit'] != item.get('unit'):
                issues.append('Quoted unit differs from estimating reference; confirm billing basis')
            reference = number(item.get('draft_quantity'))
            if reference is not None and quantity is not None and reference != quantity:
                issues.append('Quoted quantity differs from estimating reference; reconcile scope and waste')
            owner = item.get('cost_owner_row_id') or item.get('covered_by_package')
            if owner and owner != identity:
                issues.append('Estimate already assigns this scope to another cost owner: ' + owner)
            if item.get('quote_action', '').startswith('Reference heading only'):
                issues.append('Reference heading has a separate charge')
            charged[identity] = amount
        elif status != 'Included in package' and any(not blank(response.get(k)) for k in ('quoted_quantity', 'quoted_unit', 'unit_price', 'line_amount')):
            issues.append('Billing fields supplied without separately priced status')
        if status == 'Reference only' and not item.get('quote_action', '').startswith('Reference heading only'):
            issues.append('Active work marked Reference only; scope responsibility needs review')
        reviewed.append({**response, 'issues': issues, 'source_scope_issues': item.get('issues', []),
                         'package_scope': item.get('package_scope', [])})
    line_total = sum(charged.values(), Decimal(0)) if charged and all(v is not None for v in charged.values()) else None
    additional = Decimal(0)
    extras_known = True
    for category in ('delivery', 'tax'):
        inclusion = metadata.get(category + '_included')
        inclusion = inclusion.strip().lower() if isinstance(inclusion, str) else None
        try:
            raw = metadata.get(category + '_charge')
            amount = None if blank(raw) else number(raw)
        except ValueError:
            amount = None
            findings.append('Invalid ' + category + ' charge')
        if inclusion == 'yes':
            if amount not in (None, Decimal(0)):
                findings.append(category.title() + ' is included but also charged separately')
                extras_known = False
        elif inclusion == 'no' and amount is not None:
            additional += amount
        else:
            findings.append('Confirm ' + category + ' inclusion and any separate charge')
            extras_known = False
    if not blank(metadata.get('other_charges_exclusions')) or not blank(metadata.get('additional_work_options')):
        findings.append('Narrative charges, exclusions or options require review before total reconciliation')
        extras_known = False
    try:
        quoted = None if blank(metadata.get('quoted_total')) else number(metadata['quoted_total'])
    except ValueError:
        quoted = None
    if quoted is None:
        findings.append('Quoted total is missing or invalid')
    compared = line_total + additional if line_total is not None and extras_known else None
    if compared is not None and quoted is not None and compared != quoted:
        findings.append('Separate line amounts plus stated delivery/tax differ from quoted total')
    return {'snapshot_sha256': current_snapshot, 'plan_sha256': scope['plan_sha256'],
            'trade': scope['trade'], 'quote': metadata, 'responses': reviewed, 'issues': findings,
            'unanswered_scope_ids': [r['scope_id'] for r in responses if blank(r.get('response'))],
            'charged_line_total': str(line_total) if line_total is not None else None,
            'total_with_explicit_charges': str(compared) if compared is not None else None,
            'pricing_approved': False, 'estimate_changed': False, 'purchase_authorized': False,
            'remaining': 'Review supplier identity/date, product and labor inclusions, billing basis, unresolved source scope and any exclusions before accepting pricing.'}


def revision_basis(scope, issued, workbook_identity, current_snapshot):
    if scope.get('snapshot_sha256') != current_snapshot:
        raise ValueError('Current trade scope does not match the current estimate snapshot')
    expected = [issued['plan_sha256'], issued['snapshot_sha256'], issued['measurement_version']]
    if workbook_identity != expected:
        raise ValueError('Workbook drawing or estimate revision does not match issued scope')
    revision_fields = {'snapshot_sha256', 'measurement_version'}
    if ({k:v for k,v in scope.items() if k not in revision_fields}
            != {k:v for k,v in issued.items() if k not in revision_fields}):
        raise ValueError('Trade scope changed; reconcile the returned bid against the new scope')
    return {'issued_snapshot_sha256': issued['snapshot_sha256'],
            'current_snapshot_sha256': current_snapshot,
            'issued_measurement_version': issued['measurement_version'],
            'current_measurement_version': scope['measurement_version'],
            'trade_content_unchanged': True}


def read_response(scope_path, workbook_path, current_snapshot, issued_scope_path=None):
    scope_path, workbook_path = Path(scope_path), Path(workbook_path)
    scope_bytes, workbook_bytes = scope_path.read_bytes(), workbook_path.read_bytes()
    scope = json.loads(scope_bytes)
    issued_path = Path(issued_scope_path) if issued_scope_path else scope_path
    issued_bytes = issued_path.read_bytes()
    issued = json.loads(issued_bytes)
    with workbook_path.open('rb') as stream:
        workbook = load_workbook(stream, data_only=False, read_only=True)
        try:
            sheet, details = workbook['Response'], workbook['Scope details']
            if [sheet.cell(16, c).value for c in range(1, 12)] != HEADERS:
                raise ValueError('Response table headers changed')
            revision = revision_basis(scope, issued,
                [details['C5'].value, details['C6'].value, details['C7'].value], current_snapshot)
            def value(cell):
                if cell.data_type in ('f', 'e'):
                    raise ValueError('Response formulas/errors require explicit value review: ' + cell.coordinate)
                v = cell.value
                return v.isoformat() if isinstance(v, (date, datetime)) else v
            metadata = {k: value(sheet[addr]) for k, addr in METADATA.items()}
            expected = {i['row_id']: i for i in scope['items']}; responses = []
            for cells in sheet.iter_rows(min_row=17, max_col=11):
                row = [value(c) for c in cells]
                if all(blank(v) for v in row):
                    continue
                item = expected.get(row[0])
                if item is None or row[:4] != [item['row_id'], item['name'], item.get('draft_quantity'), item.get('unit')]:
                    raise ValueError('Returned scope ID or reference fields changed')
                responses.append(dict(zip(['scope_id','name','reference_quantity','reference_unit',
                    'response','included_under','quoted_quantity','quoted_unit','unit_price','line_amount','notes'], row)))
        finally:
            workbook.close()
    result = review_response(scope, metadata, responses, current_snapshot)
    if (scope_path.read_bytes() != scope_bytes or workbook_path.read_bytes() != workbook_bytes
            or issued_path.read_bytes() != issued_bytes):
        raise ValueError('Response source changed during review')
    result['source_hashes'] = {str(scope_path): hashlib.sha256(scope_bytes).hexdigest(),
                              str(issued_path): hashlib.sha256(issued_bytes).hexdigest(),
                              str(workbook_path): hashlib.sha256(workbook_bytes).hexdigest()}
    result['revision_review'] = revision
    return result


def catalog_scope(catalog, workspace, identity, scope_ids, current_snapshot):
    """Resolve a returned form by its drawing/revision and complete scope IDs."""
    if catalog['snapshot_sha256'] != current_snapshot:
        raise ValueError('Response catalog differs from the current estimate snapshot')
    if any(not isinstance(value, str) or not value for value in scope_ids) or len(set(scope_ids)) != len(scope_ids):
        raise ValueError('Returned form needs unique nonblank scope IDs')
    root = Path(workspace).resolve()
    def local(name):
        path = (root / name).resolve()
        if not path.is_relative_to(root):
            raise ValueError('Response catalog path leaves the workspace')
        return path
    matches = []
    for entry in catalog['forms']:
        issued_path = local(entry['issued_scope'])
        issued = json.loads(issued_path.read_bytes())
        expected = [issued['plan_sha256'], issued['snapshot_sha256'], issued['measurement_version']]
        if identity == expected and sorted(scope_ids) == sorted(i['row_id'] for i in issued['items']):
            baseline = local(entry['workbook'])
            if hashlib.sha256(baseline.read_bytes()).hexdigest() != entry['sha256']:
                raise ValueError('Preserved issued form changed; renew its catalog review')
            matches.append((local(entry['current_scope']), issued_path))
    if len(matches) != 1:
        raise ValueError('Returned form has no unique current catalog match; reconcile its revision and scope IDs')
    return matches[0]


def read_catalog_response(catalog_path, workspace, workbook_path, current_snapshot):
    """Use the current form catalog without inferring a trade from the filename."""
    catalog_path, workbook_path = Path(catalog_path), Path(workbook_path)
    catalog_bytes, returned_bytes = catalog_path.read_bytes(), workbook_path.read_bytes()
    with workbook_path.open('rb') as stream:
        workbook = load_workbook(stream, data_only=False, read_only=True)
        try:
            details, response = workbook['Scope details'], workbook['Response']
            identity = [details[f'C{n}'].value for n in (5, 6, 7)]
            scope_ids = [row[0].value for row in response.iter_rows(min_row=17, max_col=11)
                         if any(not blank(c.value) for c in row)]
        finally:
            workbook.close()
    scope, issued = catalog_scope(json.loads(catalog_bytes), workspace, identity, scope_ids, current_snapshot)
    result = read_response(scope, workbook_path, current_snapshot, issued)
    if catalog_path.read_bytes() != catalog_bytes or workbook_path.read_bytes() != returned_bytes:
        raise ValueError('Response catalog or returned workbook changed during review')
    result['source_hashes'][str(catalog_path)] = hashlib.sha256(catalog_bytes).hexdigest()
    result['catalog_match'] = {'current_scope': str(scope), 'issued_scope': str(issued)}
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument('--scope')
    source.add_argument('--catalog', help='Current form catalog, supporting preserved older trade forms')
    parser.add_argument('--workspace', default='.', help='Root for catalog-relative paths')
    parser.add_argument('--workbook', required=True)
    parser.add_argument('--current-snapshot', required=True)
    parser.add_argument('--issued-scope', help='Preserved original trade JSON; required for an older response revision')
    args = parser.parse_args()
    if args.catalog:
        if args.issued_scope:
            parser.error('--issued-scope is resolved from --catalog')
        result = read_catalog_response(args.catalog, args.workspace, args.workbook, args.current_snapshot)
    else:
        result = read_response(args.scope, args.workbook, args.current_snapshot, args.issued_scope)
    print(json.dumps(result, indent=2, allow_nan=False))
