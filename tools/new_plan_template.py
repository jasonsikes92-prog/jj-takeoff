"""Initialize a new job from template identities, never another job's scope or prices."""
import hashlib
import re
from pathlib import Path
import openpyxl

DEFAULT_TEMPLATE=Path(__file__).resolve().parents[1]/'templates/estimate-template.xlsx'


def read_template(path=DEFAULT_TEMPLATE):
    path=Path(path)
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    workbook=openpyxl.load_workbook(path,read_only=True,data_only=False)
    try:
        sheet=workbook['Estimate']
        columns={'name':0,'parent':1,'cost_type':2,'unit':8,'markup_pct':10}
        headers=next(sheet.iter_rows(min_row=1,max_row=1,max_col=15,values_only=True))
        for index,label in ((0,'Name'),(1,'Parent Name'),(2,'Cost Type'),(8,'Unit'),(10,'Markup %')):
            if headers[index]!=label:raise ValueError('Estimate template columns changed')
        rows=[]
        for index,values in enumerate(sheet.iter_rows(min_row=2,max_col=15,values_only=True),2):
            if values[0] is None or not str(values[0]).strip():continue
            row={key:str(values[col]) if values[col] is not None else '' for key,col in columns.items()}
            row.update(row_id=f'T-{digest[:12]}-R{index:04}',excel_row=str(index),
                completion_status='rollup_only' if row['cost_type']=='GROUP' else 'not_yet_reconciled',
                evidence=[],current_price_certified=False,whole_house_released=False)
            rows.append(row)
    finally:
        workbook.close()
    if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:raise ValueError('Template changed during intake')
    return {'template_sha256':digest,'template_source_kind':'original_workbook','rows':rows}


def initial_rules(plan_sha256,measurements,roof_ids,template,electrical_ids=(),area_ids=()):
    """Propose source candidates as partial inputs; scope review still required."""
    by_id={m['id']:m for m in measurements}
    if len(roof_ids)!=len(set(roof_ids)) or any(i not in by_id for i in roof_ids):
        raise ValueError('Roof mapping requires unique extracted measurements')
    rules=[]
    if roof_ids:
        names=('Shingle Roof - Turnkey','Metal Roofing - Exposed Fastener','Metal Roofing - Standing Seam')
        targets=[]
        for name in names:
            rows=[r for r in template['rows'] if r['name'].strip()==name and r['parent']=='Roofing']
            if len(rows)!=1:raise ValueError('Roofing template row missing or ambiguous: '+name)
            targets.append(rows[0]['excel_row'])
        rules.append({'id':'candidate-roof-surfaces','label':'Roof surface candidates; roofing material allocation unresolved',
            'measurement_ids':list(roof_ids),'template_rows':targets,'unit':'SF','rounding':'none',
            'use':'assembly_input','defer_pending_scope':True,
            'basis':'Closed roof CAD faces with embedded or explicitly reviewed source-sheet pitch labels; current face and scale review required',
            'remaining':['Verify complete roof coverage and resolve overlapping faces.',
                'Select roofing material for each face; these alternative rows share one reference, not three purchases.',
                'Calculate material waste, accessories and complete installed scope before pricing.']})
        for identity in roof_ids:
            by_id[identity]['dependent_rows']=list(targets)
    if electrical_ids:
        if len(electrical_ids)!=len(set(electrical_ids)) or any(i not in by_id or by_id[i]['kind']!='count' for i in electrical_ids):
            raise ValueError('Electrical mapping requires unique count candidates')
        targets=[r['excel_row'] for r in template['rows'] if r['name']=='Electrical - Quote' and r['parent']=='Electrical']
        if len(targets)!=1:raise ValueError('Electrical package row missing or ambiguous')
        for identity in electrical_ids:
            by_id[identity]['dependent_rows']=list(targets)
            rules.append({'id':'candidate-electrical-'+identity,'label':by_id[identity]['label'],
                'measurement_ids':[identity],'template_rows':list(targets),'unit':'EA','rounding':'whole_up',
                'use':'assembly_input','defer_pending_scope':True,
                'basis':'Located electrical symbol candidates from source-bound definitions; physical meaning and complete coverage require review.',
                'remaining':['Review each symbol group, exclude duplicate detections and resolve redlines.',
                    'Counts are partial package inputs, not additional contract charges, circuit design or code certification.',
                    'Confirm current quote and included work before pricing.']})
    if area_ids:
        if len(area_ids)!=len(set(area_ids)) or any(i not in by_id or by_id[i]['kind']!='area' for i in area_ids):
            raise ValueError('Building area mapping requires unique area measurements')
        names={'BASEMENT':'SF BASEMENT','FIRST FLOOR':'SF FIRST FLOOR','MAIN FLOOR':'SF FIRST FLOOR',
            'SECOND FLOOR':'SF SECOND FLOOR','THIRD FLOOR':'SF THIRD FLOOR','GARAGE':'SF GARAGE',
            'COVERED PORCH':'Covered Porches'}
        grouped={}
        for identity in area_ids:
            measurement=by_id[identity];label=measurement.get('source_area_label')
            target_name=names.get(label)
            if label in ('HEATED','HEATED AREA','LIVING AREA'):
                target_name=names.get(measurement.get('source_story_caption',{}).get('story'))
            if isinstance(label,str) and re.fullmatch(r'(?:(?:FRONT|REAR|BACK|SIDE) )?COVERED (?:PORCH|PATIO)(?: (?:LEFT|RIGHT))?',label):
                target_name='Covered Porches'
            if target_name is None:
                measurement['scope_mapping_reason']=('Drawing label does not identify a template story or covered-porch scope. '
                    'Keep as a gross area reference; do not infer floor level, roof coverage or finished flooring.')
                continue
            grouped.setdefault(target_name,[]).append(identity)
        for name,ids in grouped.items():
            targets=[r['excel_row'] for r in template['rows'] if r['name']==name and r['parent'].strip().upper()=='INPUTS']
            if len(targets)!=1:raise ValueError('Building area template input missing or ambiguous: '+name)
            rules.append({'id':'candidate-building-area-'+targets[0],'label':name+' - gross drawing area',
                'measurement_ids':ids,'template_rows':targets,'unit':'SF','rounding':'none',
                'use':'template_quantity','requires_geometry_review':True,'defer_pending_scope':True,
                'disjoint_areas':True,'basis':'Native area label and polygon identify a proposed building input; review extent, scale and complete coverage.',
                'remaining':['Gross building reference only; no flooring, slab, insulation or framing purchase quantity is implied.',
                    'Verify the story and complete scope; do not treat an unlabeled patio as covered porch or include it in total framed area.']})
            for identity in ids:by_id[identity]['dependent_rows']=list(targets)
        included=[identity for ids in grouped.values() for identity in ids]
        unresolved=[identity for identity in area_ids if identity not in included and not
            re.fullmatch(r'UNCOVERED PATIO(?: (?:LEFT|RIGHT))?',by_id[identity].get('source_area_label',''))]
        if 'SF FIRST FLOOR' in grouped and not unresolved:
            targets=[r['excel_row'] for r in template['rows'] if r['name']=='SF TOTAL FRAMED'
                     and r['parent'].strip().upper()=='INPUTS']
            if len(targets)!=1:raise ValueError('Total framed template input missing or ambiguous')
            rules.append({'id':'candidate-building-area-total','label':'Total framed area - confirm complete building coverage',
                'measurement_ids':included,'template_rows':targets,'unit':'SF','rounding':'none',
                'use':'template_quantity','requires_geometry_review':True,'defer_pending_scope':True,
                'disjoint_area_components':True,
                'surface_components':[{'id':identity,'measurement_id':identity,'kind':'area','operation':'add'} for identity in included],
                'basis':'Sum each measured story, garage and covered porch once. Separate scope approval must confirm every applicable building area is present.',
                'remaining':['Confirm complete coverage, including any missing floors, garages and covered spaces before approving this total.',
                    'This is gross building area, not framing labor billable layers. Determine porch/deck construction and the contractor billing basis separately.',
                    'Uncovered patios and printed schedule totals are excluded; no material quantity or price is implied.']})
            for identity in included:
                by_id[identity]['dependent_rows']=list(dict.fromkeys(by_id[identity]['dependent_rows']+targets))
    return {'plan_sha256':plan_sha256,'rules':rules}
