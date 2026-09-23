"""Read returned owner answers and retain their evidence without silently applying decisions."""
import argparse
import hashlib
import io
import json
from datetime import date, datetime, timezone
from pathlib import Path
import openpyxl

SECTIONS=(('Roberts questions','roberts','project'),('Later business decisions','later','business'))
HEADERS=['ID','Topic','Question','Context / choices','Your answer']
LEDGE_HEADERS=['ID','Topic','Question','Already saved / context','Jason’s answer']
UNKNOWN={'unknown','not sure','unsure',"don't know",'dont know','tbd'}


def collect(catalog, sheets):
    """Validate question identity even if complete rows were sorted in Excel."""
    sections=[s for s in SECTIONS if s[1]=='roberts' or catalog[s[1]]]
    expected_names={s[0] for s in sections}
    if set(sheets)!=expected_names:raise ValueError('Confirmation sheets differ from the question catalog')
    answers=[];missing=[];seen=set()
    for name,section,scope in sections:
        if any(len(q) not in (4,5) for q in catalog[section]):
            raise ValueError('Questions require four fields and an optional source reference')
        expected={q[0]:q for q in catalog[section]}
        if len(expected)!=len(catalog[section]) or seen.intersection(expected):
            raise ValueError('Unique question IDs required')
        sheet=sheets[name]
        if sheet['headers'] not in (HEADERS,LEDGE_HEADERS):raise ValueError('Confirmation column headings changed: '+name)
        found=set()
        for row in sheet['rows']:
            values=row['values'];identity=values[0]
            if all(v is None or (isinstance(v,str) and not v.strip()) for v in values):continue
            if identity not in expected or identity in found:
                raise ValueError('Unknown, missing or duplicate question ID at '+name+'!A'+str(row['row']))
            found.add(identity);question=expected[identity]
            if values[:4]!=question[:4]:raise ValueError('Question or context changed: '+identity)
            value=values[4]
            if row['answer_type'] in ('f','e'):
                raise ValueError('Enter a direct answer, not a formula/error: '+identity)
            if value is None or (isinstance(value,str) and not value.strip()):
                missing.append(identity);continue
            if isinstance(value,(date,datetime)):value=value.isoformat()
            if not isinstance(value,(str,int,float,bool)):
                raise ValueError('Unsupported answer value: '+identity)
            unknown=isinstance(value,str) and value.strip().casefold() in UNKNOWN
            answers.append({'question_id':identity,'topic':question[1],'question':question[2],
                'context':question[3],'source_reference':question[4] if len(question)==5 else None,'scope':scope,
                'sheet':name,'cell':'E'+str(row['row']),'answer':value,
                'answer_cell_type':row['answer_type'],'answer_number_format':row['number_format'],
                'status':'unresolved' if unknown else 'reported_pending_application',
                'applied_to_estimate':False,'applied_to_company_profile':False})
        if found!=set(expected):raise ValueError('Question rows missing from '+name)
        seen.update(found)
    return {'project':catalog['project'],'answers':answers,'unanswered_question_ids':missing,
        'question_count':len(seen),'answered_count':len(answers),
        'unresolved_answer_count':sum(a['status']=='unresolved' for a in answers),
        'pending_application_count':sum(a['status']=='reported_pending_application' for a in answers),
        'estimate_changed':False,'company_profile_changed':False}


def read_workbook(raw, catalog):
    workbook=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    try:
        sheets={}
        for sheet in workbook:
            rows=[]
            for cells in sheet.iter_rows(min_row=6,max_col=5):
                rows.append({'row':cells[0].row if cells[0].value is not None else cells[4].row if cells[4].value is not None else 0,
                    'values':[c.value for c in cells],'answer_type':cells[4].data_type,
                    'number_format':cells[4].number_format})
            sheets[sheet.title]={'headers':[sheet.cell(5,c).value for c in range(1,6)],'rows':rows}
        return collect(catalog,sheets)
    finally:workbook.close()


def inspect(workbook_path,catalog_path,job_dir):
    workbook_path,catalog_path,job_dir=Path(workbook_path),Path(catalog_path),Path(job_dir)
    raw=workbook_path.read_bytes();catalog_raw=catalog_path.read_bytes()
    plan=(job_dir/'plan.pdf').read_bytes()
    plan_sha=hashlib.sha256(plan).hexdigest()
    if not plan.startswith(b'%PDF-') or json.loads((job_dir/'measurements.json').read_bytes())['plan_sha256']!=plan_sha:
        raise ValueError('Job drawing differs from its measurement configuration')
    result=read_workbook(raw,json.loads(catalog_raw))
    result.update(workbook_filename=workbook_path.name,workbook_sha256=hashlib.sha256(raw).hexdigest(),
        question_catalog_sha256=hashlib.sha256(catalog_raw).hexdigest(),plan_sha256=plan_sha,
        question_catalog=json.loads(catalog_raw))
    return result,raw


def retain(result,workbook_bytes,job_dir):
    """Immutable submissions preserve old answers; acceptance/application is a separate step."""
    if hashlib.sha256(workbook_bytes).hexdigest()!=result['workbook_sha256']:
        raise ValueError('Workbook bytes changed after inspection')
    parsed=read_workbook(workbook_bytes,result['question_catalog'])
    if any(result.get(key)!=value for key,value in parsed.items()):
        raise ValueError('Answer record differs from its workbook evidence')
    if result['answered_count']==0:return {'saved':False,'reason':'No answers entered','submission_id':None}
    body=json.dumps(result,sort_keys=True,allow_nan=False,separators=(',',':')).encode()
    identity=hashlib.sha256(body).hexdigest()
    folder=Path(job_dir)/'owner_answer_submissions';record=folder/(identity+'.json');evidence=folder/(identity+'.xlsx')
    folder.mkdir(exist_ok=True)
    if evidence.exists():
        if evidence.read_bytes()!=workbook_bytes:raise ValueError('Saved answer evidence changed')
    else:
        with evidence.open('xb') as stream:stream.write(workbook_bytes)
    if record.exists():
        saved=json.loads(record.read_bytes())
        if (saved['submission']!=result or saved['submission_id']!=identity
                or saved['workbook_evidence']!=evidence.name):
            raise ValueError('Saved answer record changed')
    else:
        saved={'submission_id':identity,'received_at':datetime.now(timezone.utc).isoformat(),
               'workbook_evidence':evidence.name,'submission':result}
        with record.open('x',encoding='utf-8') as stream:json.dump(saved,stream,indent=2,allow_nan=False);stream.write('\n')
    return {'saved':True,'submission_id':identity,'record':str(record.resolve()),
            'answered_count':result['answered_count'],'estimate_changed':False,'company_profile_changed':False}


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',required=True);parser.add_argument('--questions',required=True)
    parser.add_argument('--job-dir',required=True)
    parser.add_argument('--save',action='store_true',help='Retain the returned workbook and answer record; does not apply rules or prices')
    args=parser.parse_args(argv)
    result,raw=inspect(args.workbook,args.questions,args.job_dir)
    output=retain(result,raw,args.job_dir) if args.save else result
    print(json.dumps(output,indent=2,allow_nan=False));return 0


if __name__=='__main__':raise SystemExit(main())
