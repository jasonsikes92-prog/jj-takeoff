import openpyxl
ws=openpyxl.load_workbook(r"C:\Users\jason\.claude\skills\jnj-estimate-takeoff\templates\measurements-template.xlsx",data_only=True)["Measurements"]
keys=["fram","lumber","stud","drywall","sheetrock","paint","stain","slab","concrete","gravel","electric","duct","hvac","clean","grad","drive","sod","seed"]
for r in range(2,258):
    nm=ws.cell(r,1).value; w=ws.cell(r,4).value; t=ws.cell(r,5).value
    if nm and any(k in str(nm).lower() for k in keys):
        print(f"{r} | {str(nm)[:38]:<38} | waste={w} | {t}")
