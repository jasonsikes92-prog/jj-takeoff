import openpyxl
wb=openpyxl.load_workbook(r"C:\Users\jason\.claude\skills\jnj-estimate-takeoff\examples\Roberts Residence - reference estimate.xlsx")
print("SHEETS:",wb.sheetnames)
ws=wb["Estimate"] if "Estimate" in wb.sheetnames else wb[wb.sheetnames[0]]
print("dims",ws.dimensions)
# print filled leaf rows (quantity>0) with cost
cnt=0
for r in range(2,ws.max_row+1):
    C=ws.cell(r,3).value; A=ws.cell(r,1).value
    G=ws.cell(r,7).value; H=ws.cell(r,8).value; I=ws.cell(r,9).value
    K=ws.cell(r,11).value; N=ws.cell(r,14).value
    if (isinstance(G,(int,float)) and G not in (0,None)) :
        print(f"{r}|{C}|{A}|G={G}|H={H}|{I}|K={K}|N={N}")
        cnt+=1
print("filled leaf rows:",cnt)
