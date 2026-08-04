import openpyxl
wb=openpyxl.load_workbook(r"C:\Users\jason\.claude\skills\jnj-estimate-takeoff\examples\Roberts Residence - reference estimate.xlsx")
for t in ["Assumptions to Confirm","Takeoff Notes"]:
    ws=wb[t]; print(f"\n##### {t} (dims {ws.dimensions}) #####")
    for r in range(1,min(ws.max_row,40)+1):
        row=[ws.cell(r,c).value for c in range(1,min(ws.max_column,6)+1)]
        if any(v not in (None,"") for v in row):
            print(r,"|"," | ".join("" if v is None else str(v) for v in row))
