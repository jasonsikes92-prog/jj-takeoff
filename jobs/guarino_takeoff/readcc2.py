import openpyxl
ws=openpyxl.load_workbook(r"C:\Users\jason\Downloads\J & J Custom Homes, LLC - Cost Codes.xlsx",data_only=True)["Cost Codes"]
for r in range(60,140):
    code=ws.cell(r,2).value; title=ws.cell(r,3).value; grp=ws.cell(r,5).value
    if code or title:
        print(f"{str(code):<10} | {str(title):<34} | {grp}")
