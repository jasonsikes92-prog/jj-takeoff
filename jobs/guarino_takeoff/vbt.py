import openpyxl
ws=openpyxl.load_workbook(r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\Guarino Residence - Buildern Import.xlsx",data_only=True)["Sheet1"]
print(" | ".join(str(ws.cell(1,c).value) for c in range(1,11)))
for r in [2,3,7,8,30,40,55,60,90,120]:
    vals=[ws.cell(r,c).value for c in range(1,11)]
    print(f"r{r}:", " | ".join(str(v)[:20] for v in vals))
# check no blank groups, no blank codes
blanks=sum(1 for r in range(2,ws.max_row+1) if not str(ws.cell(r,9).value or "").strip())
nocode=sum(1 for r in range(2,ws.max_row+1) if not str(ws.cell(r,3).value or "").strip())
print("rows:",ws.max_row-1,"| blank groups:",blanks,"| blank codes:",nocode)
