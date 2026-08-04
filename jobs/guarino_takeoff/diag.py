import openpyxl
wb=openpyxl.load_workbook(r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\Guarino Residence - Estimate.xlsx")
ws=wb["Estimate"]
rows=[]
for r in range(2,ws.max_row+1):
    ct=str(ws.cell(r,3).value)
    if ct in {"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}:
        q=ws.cell(r,7).value or 0; h=ws.cell(r,8).value or 0; n=ws.cell(r,14).value or 0
        if n: rows.append((n,r,str(ws.cell(r,1).value),q,h))
rows.sort(reverse=True)
for n,r,nm,q,h in rows[:25]:
    print(f"N={n:>12,.0f}  r{r:<4} q={q:<9} h={h:<8} {nm[:45]}")
