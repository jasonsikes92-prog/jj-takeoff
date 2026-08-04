import openpyxl
p=r"C:\Users\jason\Downloads\Guarino Residence V3 - Estimate Items.xlsx"
wb=openpyxl.load_workbook(p,data_only=True)
print("SHEETS:",wb.sheetnames)
for sh in wb.sheetnames:
    ws=wb[sh]
    print(f"\n##### {sh}  dims={ws.dimensions} rows={ws.max_row} cols={ws.max_column}")
    # header row
    for r in range(1,3):
        vals=[ws.cell(r,c).value for c in range(1,min(ws.max_column,14)+1)]
        print(r,[str(v)[:18] if v is not None else "" for v in vals])
