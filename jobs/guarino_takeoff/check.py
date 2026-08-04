import openpyxl
wb=openpyxl.load_workbook(r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\Guarino Residence - Estimate.xlsx")
ws=wb["Estimate"]
print("=== SAMPLE REFORMATTED DESCRIPTIONS (col O) ===")
for r in (104,213,217,332,339,342,374,110,273,353):
    print(f"  r{r} {str(ws.cell(r,1).value)[:26]:<26} -> {ws.cell(r,15).value}")
print("\n=== MEASUREMENTS tab (master bath area rows) ===")
m=wb["Measurements"]
for r in range(1,m.max_row+1):
    nm=m.cell(r,2).value
    if nm and any(k in str(nm) for k in ["Master","Bath 1","Tile","shower","Drywall","Metal Roof","Metal wall","Standing","Insulation","STANDING"]):
        print(f"  {m.cell(r,1).value} | {str(nm)[:34]:<34} | calc={str(m.cell(r,3).value)[:46]:<46} | {m.cell(r,5).value} | {m.cell(r,6).value} {m.cell(r,7).value}")
print("\n=== ALLOWANCES tab (first 16 rows) ===")
al=wb["Allowances"]
for r in range(1,17):
    print("  ",[ (str(al.cell(r,c).value)[:30] if al.cell(r,c).value not in (None,"") else "") for c in range(1,6)])
print("\n=== SELECTIONS tab (first 10 rows) ===")
s=wb["Selections"]
for r in range(1,11):
    print("  ",[ (str(s.cell(r,c).value)[:18] if s.cell(r,c).value not in (None,"") else "") for c in range(1,8)])
