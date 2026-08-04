import openpyxl
wb=openpyxl.load_workbook(r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\J & J Pre-Estimate Intake Form.xlsx")
print("Tabs:",wb.sheetnames)
ws=wb["Intake Form"]; secs=[ws.cell(r,1).value for r in range(5,ws.max_row+1) if ws.cell(r,1).fill.fgColor.rgb=="FF305496"]
print("Sections:",len([s for s in secs if s]))
for s in secs:
    if s: print("  -",s)
print("Total question rows:", sum(1 for r in range(5,ws.max_row+1) if ws.cell(r,1).value and ws.cell(r,1).fill.fgColor.rgb!="FF305496"))
