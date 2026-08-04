# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SEC=PatternFill("solid",fgColor="305496"); SECF=Font(bold=True,color="FFFFFF",size=12)
SUBF=Font(bold=True); ANSF=PatternFill("solid",fgColor="FFF9D6")
thin=Side(style="thin",color="CCCCCC"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
WRAP=Alignment(wrap_text=True,vertical="top")

# (kind, a, b, c)  kind: SEC | Q ; for Q: a=question, b=guarino answer, c=options/why
ROWS=[
 ("SEC","PROJECT & LOT",""),
 ("Q","Job / Client name","Guarino Residence",""),
 ("Q","Lot address","Lot 11 River Cove Meadows, Monticello",""),
 ("Q","County","Newton","Permits swing ~2x by county — pre-qualify permit cost before contract"),
 ("Q","Lot condition","Pasture (subdivision corner lot)","wooded / pasture / cleared / existing structure — drives clearing days, silt fence, demo"),
 ("Q","Heated SF / Garage SF / Covered porch SF","2,637 / 1,204 (conditioned) / ~347","Confirm vs the AREAS schedule; note if garage is conditioned"),
 ("Q","Build duration (months)","10","Default 10"),
 ("SEC","FOUNDATION",""),
 ("Q","Foundation type per area","Monolithic slab","slab / crawl / walkout basement — CONFIRM off the foundation plan, NOT boilerplate notes"),
 ("Q","If crawl: encapsulated? + stem-wall height","n/a","Assume encapsulated; height drives block count"),
 ("Q","If mono slab: thickened-edge note","Applied","Thickened edge = EXTRA concrete + rebar; NO separate footer labor (slab labor covers it)"),
 ("Q","Termite treatment?","Yes",""),
 ("SEC","WATER / SEWER / GAS",""),
 ("Q","Water: well or public meter? (well distance if well)","Public meter — tap 8in main","Well distance drives well electrical $/LF"),
 ("Q","Sewer: septic or public? (engineered?)","New septic (~$13k, likely engineered)","Get septic quote on a new lot / 4+ bath"),
 ("Q","Gas: propane / all-electric / natural? (tank size, buried?)","Propane 1,000-gal, buried","Tank size + bury cost"),
 ("Q","Gas appliances? Water heaters (electric HP / gas tankless)?","2 gas ranges; 2 electric heat-pump WH","Drives gas lines + propane sizing"),
 ("SEC","PRE-CON SOFT COSTS  (carry? Y/N — sometimes already done)",""),
 ("Q","Architectural Plans","N (already drawn)","ARE COGS, but ask — plans may be done"),
 ("Q","Engineering & Surveys","N","ARE COGS — include if not already done"),
 ("Q","Boundary Survey","N","Site plan may call for it"),
 ("Q","Erosion Control Plan","N","County land-disturbance requirement"),
 ("Q","(Permit, temp toilet, temp utilities = always carried)","always","—"),
 ("SEC","STRUCTURE — ROOF / INSULATION / HVAC",""),
 ("Q","Roof framing","I-joist as floor (½:12 & 1:12)","conventional rafters/trusses vs I-joist-as-floor (low slope). I-joist -> price as eng floor, drop rafters"),
 ("Q","Roof finish + rate","Standing-seam metal @ $9/SF","shingle / standing-seam metal / membrane (TPO-EPDM) / other. Waste: metal 15%, shingle 25%"),
 ("Q","Insulation spec","Open-cell foam (~$1/SF), roof + walls","OPEN-CELL is J&J default (~$1/SF). closed-cell only if specified"),
 ("Q","Ceilings: drywall or exposed?","Exposed structure, painted black","Exposed -> NO ceiling drywall; add black-coat line (heated+garage SF)"),
 ("Q","HVAC: # systems? garage conditioned? zoning?","2 systems + master damper; garage conditioned","Foam can allow 1, but conditioned garage / comfort can drive 2"),
 ("Q","Exposed rafters -> hard spiral duct? HVAC quote in hand?","Yes; quote ~$68k","Exposed rafters force hard spiral duct = big HVAC premium"),
 ("SEC","CLADDING / EXTERIOR",""),
 ("Q","Cladding type(s) + approx %","Standing-seam metal (all)","Hardie lap $3.15 / B&B / vertical panel / metal / stucco $14-16 / brick / stone $25-30 / vinyl"),
 ("Q","Pre-finished or field-paint?","Pre-finished metal -> $0 exterior paint","Pre-finished (ColorPlus / metal) cuts exterior paint"),
 ("Q","Accent / feature wall (stone/masonry)?","Yes ~282 SF stone","Carry a small masonry line even on a non-masonry house"),
 ("Q","Gutters / roof drainage","Confirm (low-slope = scuppers/internal?)","K-style vs scuppers on low slope"),
 ("SEC","INTERIOR FINISH & SELECTIONS",""),
 ("Q","Finish tier","High-end","mid / upper-mid / high-end — bumps flooring, cabinets, counters, fixtures"),
 ("Q","Flooring per area","Stained concrete + tile in baths","LVP / wood / tile / stained concrete"),
 ("Q","Countertops","Level-6 granite/quartz (or concrete)",""),
 ("Q","Cabinets (tier + LF)","High-end; kitchen+butler+pantry+vanities","stock / semi-custom / custom"),
 ("Q","Appliances (tier; owner-supplied?)","Mid-high allowance incl 2 gas ranges","Often owner-supplied — flag"),
 ("Q","Interior doors (style; solid/hollow)","Flat-stock; solid at all bedrooms",""),
 ("Q","Fireplace? Garage door? Windows?","No FP; flush 16x8; black vinyl SH/FX no grid",""),
 ("SEC","SITE & LOGISTICS",""),
 ("Q","Demo required?","No","Books to Excavation (03.05) code"),
 ("Q","Clearing days + silt fence LF","Pasture ~2 days strip; silt ~750 LF","Silt fence = disturbed-area perimeter"),
 ("Q","Driveway (length, material, apron/walks)","~110 LF concrete + walks","Never under-budget the drive"),
 ("Q","Dumpsters (count)","3",""),
 ("Q","Landscaping (sod/seed/plants/irrigation)","sod + plants + seed + irrigation","Zero-trap — always budget"),
 ("SEC","MARGIN / MARKUP  (applied in Buildern)",""),
 ("Q","Per-line markup","15 / 7 / 8 by type (Material/Labor&Sub/Allowance)","Default — entered on each line"),
 ("Q","Overhead % (Buildern Summary)","20%","NOT an estimate line — Buildern summary"),
 ("Q","Contingency % (Buildern Summary)","8% (first all-metal modern)","familiar ~3 / some-new ~5 / first-of-kind ~8-10. NOT an estimate line"),
 ("SEC","ANYTHING THE PLANS DON'T SPECIFY",""),
 ("Q","Open items / notes","","List anything the redlines don't fully resolve"),
]

def build(ws, filled):
    ws.column_dimensions["A"].width=46; ws.column_dimensions["B"].width=40; ws.column_dimensions["C"].width=58
    ws.cell(1,1).value="J & J CUSTOM HOMES — PRE-ESTIMATE INTAKE FORM"
    ws.cell(1,1).font=Font(bold=True,size=14)
    ws.cell(2,1).value=("WORKED EXAMPLE — Guarino Residence" if filled else "Fill in before estimating. Ask Jason anything not 100% clear on the plans; never guess.")
    ws.cell(2,1).font=Font(italic=True,color="888888")
    ws.cell(4,1).value="Question / Item"; ws.cell(4,2).value="Answer"; ws.cell(4,3).value="Options / why it matters / default"
    for c in range(1,4): ws.cell(4,c).font=Font(bold=True); ws.cell(4,c).fill=PatternFill("solid",fgColor="D9E1F2"); ws.cell(4,c).border=BORD
    r=5
    for item in ROWS:
        if item[0]=="SEC":
            ws.cell(r,1).value=item[1]
            for c in range(1,4): ws.cell(r,c).fill=SEC; ws.cell(r,c).font=SECF; ws.cell(r,c).border=BORD
            ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
        else:
            _,q,gA,why=item
            ws.cell(r,1).value=q; ws.cell(r,1).alignment=WRAP
            ws.cell(r,2).value=(gA if filled else ""); ws.cell(r,2).fill=ANSF; ws.cell(r,2).alignment=WRAP
            ws.cell(r,3).value=why; ws.cell(r,3).alignment=WRAP; ws.cell(r,3).font=Font(size=9,color="666666")
            for c in range(1,4): ws.cell(r,c).border=BORD
        r+=1
    ws.freeze_panes="A5"

wb=openpyxl.Workbook(); build(wb.active,False); wb.active.title="Intake Form"
g=wb.create_sheet("Guarino (example)"); build(g,True)
for path in [r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\J & J Pre-Estimate Intake Form.xlsx",
             r"C:\Users\jason\.claude\skills\jnj-estimate-takeoff\templates\pre-estimate-intake-form.xlsx"]:
    wb.save(path); print("SAVED",path)
