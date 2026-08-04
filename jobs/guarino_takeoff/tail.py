import openpyxl
p=r"C:\Users\jason\Downloads\Guarino Residence V3 - Estimate Items.xlsx"
# formulas view
wf=openpyxl.load_workbook(p); wsf=wf["Estimate"]
wd=openpyxl.load_workbook(p,data_only=True); wsd=wd["Estimate"]
print("rows 682-693 all cols (formula / value):")
for r in range(682,694):
    for c in range(1,16):
        f=wsf.cell(r,c).value; v=wsd.cell(r,c).value
        if f not in (None,"") or v not in (None,""):
            from openpyxl.utils import get_column_letter
            print(f"  {get_column_letter(c)}{r}: formula={f!r} value={v!r}")
