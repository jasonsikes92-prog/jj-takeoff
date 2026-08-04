import openpyxl
wb=openpyxl.load_workbook(r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\Guarino Residence - Estimate.xlsx",data_only=False)
ws=wb["Estimate"]
LEAF={"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}
tot=0;errs=0
for r in range(2,ws.max_row+1):
    for c in range(1,16):
        v=ws.cell(r,c).value
        if isinstance(v,str) and v.startswith("#") and v.endswith(("!","?")): errs+=1
    if str(ws.cell(r,3).value) in LEAF:
        tot+=ws.cell(r,14).value or 0
# scan for any lingering O&P / contingency text
bad=[]
for r in range(2,ws.max_row+1):
    for c in (1,13):
        v=str(ws.cell(r,c).value or "")
        if any(k in v.lower() for k in ["contingency","overhead & profit","contract price","o&p"]): bad.append((r,v))
print("Estimate leaf N sum:",round(tot,2)," | error cells:",errs)
print("O&P/contingency lines remaining:",bad if bad else "NONE ✔")
print("Tabs:",wb.sheetnames)
