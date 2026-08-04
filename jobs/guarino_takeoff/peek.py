import openpyxl
p=r"C:\Users\jason\Downloads\Guarino Residence V3 - Estimate Items.xlsx"
wb=openpyxl.load_workbook(p,data_only=True); ws=wb["Estimate"]
want=["Mechanical / HVAC","Roofing","Insulation","Monolithic Slab","Foundation","Drywall","General Requirements","Plumbing","Landscaping","Siding","Masonry"]
cur=None
for r in range(2,ws.max_row+1):
    nm=ws.cell(r,1).value; ct=ws.cell(r,3).value; q=ws.cell(r,7).value; h=ws.cell(r,8).value; n=ws.cell(r,14).value
    if ct=="GROUP": cur=str(nm)
    if isinstance(n,(int,float)) and n and ct in {"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}:
        par=str(ws.cell(r,2).value)
        # show HVAC, Roofing, Insulation, Drywall, GenReq, Foundation lines
        if any(k in (cur or "") or k in par or k in str(nm) for k in ["HVAC","Mechanical","Roof","Insulation","Slab","Foundation","Footer","Drywall","Permit","Engineering","Erosion","Survey","Plan Design","Septic","Water","Stucco","Stone","Brick","Siding","Concrete"]):
            print(f"r{r:<3} q={str(q):<8} h={str(h):<8} N={n:>9,.0f}  {str(nm)[:50]}")
