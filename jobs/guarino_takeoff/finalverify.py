import openpyxl
LEAF={"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}
# Working estimate
ws=openpyxl.load_workbook("Guarino Residence - Estimate.xlsx",data_only=False)["Estimate"]
tot=0;errs=0;foot=0
for r in range(2,ws.max_row+1):
    for c in range(1,16):
        v=ws.cell(r,c).value
        if isinstance(v,str) and v.startswith("#") and v.endswith(("!","?")): errs+=1
    if str(ws.cell(r,3).value) in LEAF: tot+=ws.cell(r,14).value or 0
    if "footer labor" in str(ws.cell(r,1).value or "").lower(): foot+=1
print(f"Estimate subtotal (reload): {tot:,.2f} | error cells: {errs} | footer-labor lines: {foot}")
# Buildern import
bi=openpyxl.load_workbook("Guarino Residence - Buildern Import.xlsx",data_only=True)["Sheet1"]
nlines=bi.max_row-1; blankg=blankc=0; base=mk=0
for r in range(2,bi.max_row+1):
    if not str(bi.cell(r,9).value or "").strip(): blankg+=1
    if not str(bi.cell(r,3).value or "").strip(): blankc+=1
    base+=(bi.cell(r,5).value or 0)*(bi.cell(r,7).value or 0); mk+=bi.cell(r,8).value or 0
print(f"Buildern import: {nlines} lines | blank groups: {blankg} | blank codes: {blankc}")
print(f"  base {base:,.0f} + markup {mk:,.0f} = {base+mk:,.0f}  (matches estimate: {abs((base+mk)-tot)<2})")
print("Tabs:",openpyxl.load_workbook('Guarino Residence - Estimate.xlsx').sheetnames)
