import openpyxl
p=r"C:\Users\jason\Downloads\Guarino Residence V3 - Estimate Items.xlsx"
wb=openpyxl.load_workbook(p,data_only=True); ws=wb["Estimate"]
LEAF={"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}
# name->row for parent walk
name_rows={}
for r in range(2,ws.max_row+1):
    nm=ws.cell(r,1).value
    if nm: name_rows.setdefault(str(nm),r)
def toplevel(r):
    cur=r
    for _ in range(30):
        par=ws.cell(cur,2).value
        if par is None or str(par).strip()=="": return str(ws.cell(cur,1).value)
        pr=name_rows.get(str(par))
        if pr is None or pr==cur: return str(par)
        cur=pr
    return str(ws.cell(r,1).value)
roll={}; tot=0; rows=[]
for r in range(2,ws.max_row+1):
    if str(ws.cell(r,3).value) in LEAF:
        q=ws.cell(r,7).value; h=ws.cell(r,8).value; mk=ws.cell(r,11).value; n=ws.cell(r,14).value
        if isinstance(n,(int,float)) and n:
            roll[toplevel(r)]=roll.get(toplevel(r),0)+n; tot+=n
            rows.append((n,r,str(ws.cell(r,1).value),q,h,mk))
print("JASON total Amount (sum leaf N):",round(tot,2))
# summary tail
print("\n-- tail rows (summary) --")
for r in range(682,ws.max_row+1):
    vals=[ws.cell(r,c).value for c in (1,13,14)]
    if any(v not in (None,"") for v in vals): print(r,vals)
print("\n== JASON CATEGORY ROLLUP ==")
for k,v in sorted(roll.items(),key=lambda z:-z[1]):
    print(f"  {k:<45} {v:>12,.0f}")
import pickle; pickle.dump((roll,tot,rows),open("jason.pkl","wb"))
