# -*- coding: utf-8 -*-
import openpyxl, pickle, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
OUT=r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\Guarino Residence - Estimate.xlsx"
NOTES,roll,R=pickle.load(open("results.pkl","rb"))
wb=openpyxl.load_workbook(OUT); ws=wb["Estimate"]
LEAF={"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}
HDR=PatternFill("solid",fgColor="305496"); HF=Font(bold=True,color="FFFFFF")
SUB=PatternFill("solid",fgColor="D9E1F2")

# name->row, toplevel
nr={}
for r in range(2,ws.max_row+1):
    nm=ws.cell(r,1).value
    if nm: nr.setdefault(str(nm),r)
def top(r):
    cur=r
    for _ in range(30):
        par=ws.cell(cur,2).value
        if par is None or str(par).strip()=="": return str(ws.cell(cur,1).value)
        pr=nr.get(str(par))
        if pr is None or pr==cur: return str(par)
        cur=pr
    return str(ws.cell(r,1).value)

# waste/type/plan classifier
def classify(name, code, ct):
    n=name.lower(); c=str(code)
    P_FND="Foundation/Slab (sh.4-5)"; P_RF="Roof Plan (sh.7)"; P_EL="Elevations (sh.9-10)"
    P_FP="Floor Plan (sh.6)"; P_SITE="Site Plan (sh.2)"; P_SCH="Schedules (sh.1/3/6)"
    if "metal roof" in n or ("roof" in n and "insul" not in n): return 15,"Roof area",P_RF
    if "shingle" in n: return 25,"Roof area",P_RF
    if any(k in n for k in["footer","slab","grade beam","concrete-","block","brick","flatwork","gravel"]) or c.startswith("4."): return 10,"Area/Linear",P_FND
    if "insulation" in n or "spray" in n: return 10,"Area (open-cell, vendor qty)",P_EL
    if any(k in n for k in["siding","cladding","fascia","corner","frieze","freize","porch ceiling","t&g","metal wall"]): return 10,"Area/Linear",P_EL
    if any(k in n for k in["tile","lvp","backsplash","schluter","mud bed"]): return 10,"Area",P_FP
    if "granite" in n or "quartz" in n or "countertop" in n: return 1.24,"Area",P_FP
    if any(k in n for k in["base molding","trim","crown","beam","casing"]): return 10,"Linear",P_FP
    if "drywall" in n: return 0,"Wall area (perim x ht; rate incl waste)",P_FP
    if any(k in n for k in["framing","lumber","engineered floor","electrical","paint","stain","duct","clean","grading"]): return 0,"Rate-based (waste in $/unit)",P_FP
    if any(k in n for k in["window","door","fixture","opening","cabinet","lf of","sink","toilet","fan","can light","heater","hvac","septic","meter","propane","tank","permit"]): return 0,"Count / each",P_SCH
    return 0,"Each",P_SCH

# read per-row data
rowdata={}
for r in range(2,ws.max_row+1):
    ct=ws.cell(r,3).value
    if str(ct) in LEAF:
        rowdata[r]=dict(name=str(ws.cell(r,1).value),ct=str(ct),code=ws.cell(r,5).value,
                        q=ws.cell(r,7).value or 0,unit=ws.cell(r,9).value,amt=ws.cell(r,14).value or 0,
                        desc=ws.cell(r,15).value, top=top(r))
notemap={r:basis for (r,item,q,unit,amt,basis) in NOTES}

# 1) REFORMAT DESCRIPTIONS: append qty for measured lines; counts get count; allowances keep language
for r,d in rowdata.items():
    if d["amt"]==0: continue
    waste,typ,plan=classify(d["name"],d["code"],d["ct"])
    base=d["desc"] or d["name"]
    q=d["q"]; u=(d["unit"] or "").strip()
    if d["ct"]=="ALLOWANCE":
        if not str(base).strip().endswith("allowance.") and "allowance" not in str(base).lower():
            base=str(base)+" Allowance."
        ws.cell(r,15).value=base
    else:
        qty_txt=f"{q:,.0f}" if float(q).is_integer() else f"{q:,.1f}"
        if typ.startswith("Count") or "each" in (u.lower()) or typ=="Each":
            tail=f" — {qty_txt} {u}".rstrip()
        elif waste>0:
            tail=f" — {qty_txt} {u} (incl. {waste:g}% waste)"
        else:
            tail=f" — {qty_txt} {u}"
        ws.cell(r,15).value=f"{base}{tail}"

# 2) MEASUREMENTS TAB (replace Takeoff Notes)
if "Takeoff Notes" in wb.sheetnames: del wb["Takeoff Notes"]
m=wb.create_sheet("Measurements",1)
m.append(["Cost Code","Item","Calc / basis (how derived)","Type","Waste %","Qty","Unit","Plan sheet"])
for c in range(1,9): m.cell(1,c).fill=HDR; m.cell(1,c).font=HF; m.cell(1,c).alignment=Alignment(wrap_text=True,vertical="center")
cur_top=None
for r in sorted(rowdata):
    d=rowdata[r]
    if d["amt"]==0: continue
    if d["top"]!=cur_top:
        cur_top=d["top"]; m.append([cur_top]); rr=m.max_row
        for c in range(1,9): m.cell(rr,c).fill=SUB
        m.cell(rr,1).font=Font(bold=True)
    waste,typ,plan=classify(d["name"],d["code"],d["ct"])
    calc=notemap.get(r,"") or ""
    calc=re.sub(r"\s*FLAG.*$","",calc)  # keep calc clean; flags live on Assumptions tab
    m.append([d["code"],d["name"],calc,typ,(f"{waste:g}%" if waste else "0%"),
              round(d["q"],1),d["unit"],plan])
for col,w in zip("ABCDEFGH",[9,40,60,28,8,10,9,22]): m.column_dimensions[col].width=w
m.freeze_panes="A2"
# geometry block
m.append([]); m.append(["GEOMETRY (measured)"]); m.cell(m.max_row,1).font=Font(bold=True)
geo=[("Heated SF",2637,"Sheet 1/3 AREAS table"),("Garage SF (conditioned)",1204,"Sheet 1/3"),
 ("Covered porch SF",347,"Front 77 (sch.) + rear ~270 (roof-plan backout)"),
 ("Total under roof",4188,"heated+garage+porch"),("Roof area",4850,"footprint x1.003 low-slope; +15% waste=5,578 ordered"),
 ("Exterior wall area",4682,"elevations gross (Jason measure)"),
 ("Drywall wall area",17939,"main 1,067 LF x15.5ft + garage 140 LF x10ft"),
 ("Slab / stained concrete",3841,"heated+garage"),("Foundation perimeter / footer",445,"thickened-edge LF"),
 ("Scale check","","Floor plan 106'-5\"=1427pt=13.41 pt/ft (=3/16in); elev 18pt/ft (side 54.2'=53'-9\")")]
for k,v,b in geo: m.append(["",k,b,"","",v,"",""])

# 3) ALLOWANCES TAB (by room)
al=wb.create_sheet("Allowances",2)
al.append(["Room / Area","Allowance Item","Qty","Unit","Allowance $"])
for c in range(1,6): al.cell(1,c).fill=HDR; al.cell(1,c).font=HF
from collections import defaultdict, OrderedDict
byroom=defaultdict(list)
for r in sorted(rowdata):
    d=rowdata[r]
    if d["ct"]=="ALLOWANCE" and d["amt"]:
        byroom[d["top"]].append(d)
gtot=0
for room in byroom:
    al.append([room]); rr=al.max_row
    for c in range(1,6): al.cell(rr,c).fill=SUB
    al.cell(rr,1).font=Font(bold=True)
    st=0
    for d in byroom[room]:
        al.append(["",d["name"],round(d["q"],1),d["unit"],round(d["amt"],2)]); st+=d["amt"]
    al.append(["",f"  {room} subtotal","","",round(st,2)]); al.cell(al.max_row,5).font=Font(bold=True); gtot+=st
al.append([]); al.append(["","TOTAL ALLOWANCES (incl. markup)","","",round(gtot,2)])
al.cell(al.max_row,2).font=Font(bold=True,size=12); al.cell(al.max_row,5).font=Font(bold=True,size=12)
al.cell(al.max_row,5).fill=PatternFill("solid",fgColor="FFFF00")
for col,w in zip("ABCDE",[34,46,9,9,14]): al.column_dimensions[col].width=w
al.freeze_panes="A2"
for row in al.iter_rows(min_row=2):
    if isinstance(row[4].value,(int,float)): row[4].number_format='#,##0.00'

# 4) SELECTIONS TAB
sel=wb.create_sheet("Selections",3)
sel.append(["Room / Area","Selection Item","Allowance $","Selected Product","Vendor","Actual Cost","Variance vs Allowance","Change Order?","Approved (init / date)"])
for c in range(1,10): sel.cell(1,c).fill=HDR; sel.cell(1,c).font=HF; sel.cell(1,c).alignment=Alignment(wrap_text=True,vertical="center")
for room in byroom:
    sel.append([room]); rr=sel.max_row
    for c in range(1,10): sel.cell(rr,c).fill=SUB
    sel.cell(rr,1).font=Font(bold=True)
    for d in byroom[room]:
        sel.append([ "",d["name"],round(d["amt"],2),"","","","","",""])
        rr=sel.max_row
        # variance formula = actual - allowance
        sel.cell(rr,7).value=f"=IF(F{rr}=\"\",\"\",F{rr}-C{rr})"
for col,w in zip("ABCDEFGHI",[24,40,12,24,16,12,16,12,18]): sel.column_dimensions[col].width=w
sel.freeze_panes="A2"
for row in sel.iter_rows(min_row=2):
    for c in (2,5):
        pass
    if isinstance(row[2].value,(int,float)): row[2].number_format='#,##0.00'
    row[5].number_format='#,##0.00'; row[6].number_format='#,##0.00'

# order tabs: Estimate, Measurements, Allowances, Selections, Assumptions
order=["Estimate","Measurements","Allowances","Selections"]
if "Assumptions to Confirm" in wb.sheetnames: order.append("Assumptions to Confirm")
wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
wb.save(OUT)
print("SAVED with tabs:",wb.sheetnames)
print(f"TOTAL ALLOWANCES = {gtot:,.0f}")
print("Allowance rooms:",list(byroom.keys()))
PYEOF = None
