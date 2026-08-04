import openpyxl
p=r"C:\Users\jason\Downloads\J & J Custom Homes, LLC - Cost Codes.xlsx"
wb=openpyxl.load_workbook(p,data_only=True)
print("SHEETS:",wb.sheetnames)
for sh in wb.sheetnames:
    ws=wb[sh]
    print(f"\n##### {sh}  dims={ws.dimensions} rows={ws.max_row} cols={ws.max_column}")
    for r in range(1,min(ws.max_row,60)+1):
        vals=[ws.cell(r,c).value for c in range(1,ws.max_column+1)]
        if any(v not in (None,"") for v in vals):
            print(r,"|"," | ".join("" if v is None else str(v)[:26] for v in vals))
