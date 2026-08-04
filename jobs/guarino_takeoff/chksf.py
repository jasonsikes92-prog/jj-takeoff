import openpyxl
ws=openpyxl.load_workbook("Guarino Residence - Estimate.xlsx",data_only=True)["Estimate"]
for r in range(ws.max_row-12,ws.max_row+1):
    a=ws.cell(r,13).value; b=ws.cell(r,14).value
    if a and ("SF" in str(a) or "SUBTOTAL" in str(a) or "cost" in str(a)): print(f"  {a}: {b}")
