import openpyxl, pickle
OUT=r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\Guarino Residence - Estimate.xlsx"
wb=openpyxl.load_workbook(OUT, data_only=False)
ws=wb["Estimate"]
LEAF={"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}
tot=0; errs=0; ncount=0
for r in range(2,ws.max_row+1):
    for c in range(1,16):
        v=ws.cell(r,c).value
        if isinstance(v,str) and v.startswith("#") and v.endswith(("!","?")): errs+=1
    if str(ws.cell(r,3).value) in LEAF:
        n=ws.cell(r,14).value or 0
        if isinstance(n,(int,float)): tot+=n; 
        if n: ncount+=1
print("RELOADED leaf N sum:",round(tot,2),"  filled leaf rows:",ncount)
print("Excel error cells:",errs)
NOTES,roll,R=pickle.load(open("results.pkl","rb"))
print("in-memory DIRECT SUBTOTAL:",round(R['sub'],2))
print("MATCH:" , abs(tot-R['sub'])<1.0)
print("\nSHEETS:",wb.sheetnames)
print("\n==== CATEGORY ROLLUP (direct, incl baked markup) ====")
for k,v in sorted(roll.items(),key=lambda z:-z[1]):
    print(f"  {k:<42} {v:>12,.0f}")
print(f"  {'-'*42} {'-'*12}")
print(f"  {'DIRECT SUBTOTAL':<42} {R['sub']:>12,.0f}")
print(f"  {'Contingency 8%':<42} {R['cont']:>12,.0f}")
print(f"  {'Overhead & Profit 20%':<42} {R['onp']:>12,.0f}")
print(f"  {'CONTRACT PRICE':<42} {R['sell']:>12,.0f}")
