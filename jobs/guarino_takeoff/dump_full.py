import openpyxl
wb=openpyxl.load_workbook(r"C:\Users\jason\.claude\skills\jnj-estimate-takeoff\templates\estimate-template.xlsx")
ws=wb["Estimate"]
for r in range(1,ws.max_row+1):
    A=ws.cell(r,1).value; B=ws.cell(r,2).value; C=ws.cell(r,3).value
    E=ws.cell(r,5).value; G=ws.cell(r,7).value; H=ws.cell(r,8).value; I=ws.cell(r,9).value
    K=ws.cell(r,11).value
    if A is None and B is None and H is None: continue
    print(f"{r}|{C}|A={A}|par={B}|E={E}|G={G}|H={H}|I={I}|K={K}")
