# Maps Buildern import line items to bid-package scope numbers; writes scope_lines.json.
import openpyxl, json
wb = openpyxl.load_workbook('Guarino Residence - Buildern Import.xlsx', data_only=True)
ws = wb['Sheet1']

def clean(desc):
    if not desc: return ''
    return str(desc).split('�')[0].strip()

def scope_for(name, group):
    n = name.lower()
    g = group
    if g in ('BUSINESS OPERATIONS','PRELIMINARY WORKS'): return None
    if g == 'LAND & SITE IMPROVEMENT':
        if 'erosion control' in n or 'silt fence' in n: return '01B'
        return '01'
    if g == 'PLUMBING':
        if 'septic' in n or 'sewer' in n: return '02'
        return '09'
    if g == 'HVAC':
        if 'propane tank' in n or 'bury propane' in n: return '04'
        return '10'
    if g == 'FOUNDATION & BASEMENT': return '05'
    if g == 'FRAMING':
        if any(k in n for k in ['siding','fascia','porch ceiling']): return '14'
        if 'engineered floor' in n: return '07C'
        if 'framing labor' in n: return '07B'
        if 'framing lumber' in n: return '07A'
        return '07A'
    if g == 'ROOFING': return '08'
    if g == 'WINDOWS & DOORS':
        if 'garage door' in n: return '24'
        if any(k in n for k in ["hinged","8'0\" single","8' garage entry","door knobs","door stops","interior door hardware"]): return '16'
        return '15'
    if g == 'CARPENTRY & COUNTERTOPS':
        if any(k in n for k in ['base molding','casing','trim labor','trim - labor','quarter round']): return '16'
        if 'granite or quartz' in n or 'charge per cutout' in n: return '20'
        return '19'
    if g == 'ELECTRICAL': return '11'
    if g == 'INSULATION & DRYWALL':
        if 'drywall' in n: return '13'
        return '12'
    if g == 'MASONRY & FIREPLACE': return '22'
    if g == 'FLOORING & TILE':
        if 'concrete floor' in n or 'stained' in n: return '21'
        return '18'
    if g in ('APPLIANCES','FINISHES'): return None
    if g == 'GLASS & MIRRORS': return '23'
    if g == 'PAINT & WALLCOVERING': return '17'
    if g == 'LANDSCAPING': return '27'
    if g == 'EXTERIOR WORKS': return '26B'
    if g == 'CLEANUP':
        if 'final clean interior' in n: return '28A'
        if 'window cleaning' in n: return '28B'
        if 'pressure washing' in n: return '28C'
        return None
    return None

scopes={}
for r in range(2, ws.max_row+1):
    name=ws.cell(r,1).value
    qty=ws.cell(r,5).value
    unit=ws.cell(r,6).value
    grp=ws.cell(r,9).value
    desc=clean(ws.cell(r,10).value)
    s=scope_for(str(name), grp)
    if not s: continue
    is_allow='allowance' in desc.lower()
    scopes.setdefault(s,[]).append({'name':str(name).strip(),'ct':ws.cell(r,4).value,
        'qty':qty,'unit':unit,'desc':desc,'allowance':is_allow})

with open('scope_lines.json','w',encoding='utf-8') as f:
    json.dump(scopes,f,indent=1,ensure_ascii=False)
for s in sorted(scopes):
    print(s, len(scopes[s]),'lines')
