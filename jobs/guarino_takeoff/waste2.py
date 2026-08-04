import openpyxl
p=r"C:\Users\jason\.claude\skills\jnj-estimate-takeoff\templates\measurements-template.xlsx"
ws=openpyxl.load_workbook(p,data_only=True)["Measurements"]
import re
for r in range(60,258):
    nm=ws.cell(r,1).value; q=ws.cell(r,2).value; u=ws.cell(r,3).value; w=ws.cell(r,4).value; t=ws.cell(r,5).value
    if nm and (w not in (None,"") or (t in ("Area","Linear","Roof area") )) :
        print(f"{r} | {str(nm)[:34]:<34} | waste={w} | {t}")
    elif nm and str(nm).strip() and q is None and u is None:  # section headers
        print(f"--- {nm}")
