import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\jason\.claude\skills\jnj-estimate-takeoff\templates\estimate-template.xlsx", data_only=False)
for ws in wb.worksheets:
    print(f"\n##### SHEET: {ws.title}  dims={ws.dimensions} maxrow={ws.max_row} maxcol={ws.max_column}")
    # header row
    for r in range(1, min(ws.max_row,4)+1):
        vals=[]
        for c in range(1, min(ws.max_column,16)+1):
            v=ws.cell(r,c).value
            vals.append(f"{openpyxl.utils.get_column_letter(c)}={v}")
        print(r, " | ".join(str(x) for x in vals if x and 'None' not in str(x)))
