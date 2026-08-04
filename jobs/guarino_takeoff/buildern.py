# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
EST=r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\Guarino Residence - Estimate.xlsx"
CC=r"C:\Users\jason\Downloads\J & J Custom Homes, LLC - Cost Codes.xlsx"
OUT=r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\Guarino Residence - Buildern Import.xlsx"
LEAF={"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}

# --- cost code lookup ---
ccws=openpyxl.load_workbook(CC,data_only=True)["Cost Codes"]
GROUP_BY_MAJOR={1:"PRELIMINARY WORKS",2:"BUSINESS OPERATIONS",3:"LAND & SITE IMPROVEMENT",
4:"FOUNDATION & BASEMENT",5:"FRAMING",6:"ROOFING",7:"PLUMBING",8:"ELECTRICAL",9:"HVAC",
10:"INSULATION & DRYWALL",11:"GLASS & MIRRORS",12:"CARPENTRY & COUNTERTOPS",13:"WINDOWS & DOORS",
14:"PAINT & WALLCOVERING",15:"FLOORING & TILE",16:"APPLIANCES",17:"MASONRY & FIREPLACE",
18:"EXTERIOR WORKS",19:"LANDSCAPING",20:"FINISHES",21:"CLEANUP"}
lookup={}
for r in range(2,ccws.max_row+1):
    code=ccws.cell(r,2).value; title=ccws.cell(r,3).value; grp=ccws.cell(r,5).value
    if code is None or str(code).strip()=="": continue
    try: key=round(float(str(code)),2)
    except: continue
    major=int(float(str(code)))
    g=grp if grp not in (None,"None","") else GROUP_BY_MAJOR.get(major,"")
    lookup[key]=(str(code), str(title) if title else "", g)

def map_code(raw):
    if raw is None: return ("","","")
    s=str(raw).strip()
    try: key=round(float(s),2)
    except: return (s,"","")
    if key in lookup: return lookup[key]
    major=int(float(s))
    return (s,"",GROUP_BY_MAJOR.get(major,""))

MK={"MATERIAL":0.15,"FEE":0.15,"LABOR":0.07,"SUBCONTRACTOR":0.07,"EQUIPMENT":0.07,"ALLOWANCE":0.08}
AREA_KW=["tile","insulation","drywall","siding","cladding","metal roof","roofing","paint","stain",
"backsplash","flooring","lvp","mud bed","schluter","slab","gravel","porch ceiling","stucco","stone",
"brick","coat","stained concrete","wall area"]
def fix_unit(name,unit):
    u=str(unit or "").strip(); ul=u.lower(); n=name.lower()
    if ul in ("yd","yd3","cu yd","cuyd"): return "CY"
    if ul in ("sq ft","ft2","sqft","sf"): return "sqft"
    if ul in ("feet","ft","linear foot","linear ft","lf"): return "LF"
    if ul in ("each","unit","ea") and any(k in n for k in AREA_KW) and "opening" not in n and "fixture" not in n:
        return "sqft"
    return u if u else "each"

est=openpyxl.load_workbook(EST,data_only=True)["Estimate"]
rows=[]; unmapped=[]; base_tot=0; mk_tot=0
for r in range(2,est.max_row+1):
    ct=str(est.cell(r,3).value)
    if ct not in LEAF: continue
    q=est.cell(r,7).value or 0; h=est.cell(r,8).value or 0; amt=est.cell(r,14).value or 0
    if not amt: continue
    name=str(est.cell(r,1).value); code=est.cell(r,5).value; unit=est.cell(r,9).value; desc=est.cell(r,15).value
    bcode,btitle,bgroup=map_code(code)
    if not bgroup:
        nl=name.lower()
        if any(k in nl for k in ["grad","driveway","drive","backfill","equipment","culvert"]): bgroup="EXTERIOR WORKS"
        elif any(k in nl for k in ["blueprint","plan","permit","survey","erosion","toilet","utilit"]): bgroup="PRELIMINARY WORKS"
        else: bgroup="LAND & SITE IMPROVEMENT"
    if not bgroup: unmapped.append((name,code))
    mkpct=MK.get(ct,0.0); base=q*h; mk=round(base*mkpct,2)
    base_tot+=base; mk_tot+=mk
    rows.append([name, ct.capitalize(), bcode, btitle, round(q,2), fix_unit(name,unit), round(h,4), mk, bgroup, desc or ""])

wb=openpyxl.Workbook(); ws=wb.active; ws.title="Sheet1"
hdr=["Name","Cost type","Cost code","Cost title","Quantity","Unit","Unit cost","Markup","Group","Description"]
ws.append(hdr)
for c in range(1,11): ws.cell(1,c).font=Font(bold=True)
for row in rows: ws.append(row)
# cost code as TEXT to preserve leading zeros
for r in range(2,ws.max_row+1):
    ws.cell(r,3).number_format='@'
    ws.cell(r,3).value=str(ws.cell(r,3).value)
    ws.cell(r,7).number_format='#,##0.00'; ws.cell(r,8).number_format='#,##0.00'
for col,w in zip("ABCDEFGHIJ",[42,14,10,28,10,8,11,11,26,60]): ws.column_dimensions[col].width=w
ws.freeze_panes="A2"
wb.save(OUT)
print("SAVED",OUT)
print(f"line items: {len(rows)}")
print(f"base cost (qty x unit cost): {base_tot:,.0f}")
print(f"per-line markup (15/7/8): {mk_tot:,.0f}  ({mk_tot/base_tot*100:.1f}% blended)")
print(f"= line subtotal w/ markup: {base_tot+mk_tot:,.0f}  (then +20% overhead in Buildern = {(base_tot+mk_tot)*1.20:,.0f})")
if unmapped:
    print("\nUNMAPPED group (check codes):")
    for n,c in unmapped[:20]: print(f"  code {c}: {n[:40]}")
else:
    print("\nAll lines mapped to a group. ✔")
# group rollup
from collections import defaultdict
gr=defaultdict(float)
for row in rows: gr[row[8]]+= (row[4]*row[6]+row[7])
print("\nGROUP rollup (w/ per-line markup):")
for g,v in sorted(gr.items(),key=lambda z:-z[1]): print(f"  {g:<28} {v:>11,.0f}")
