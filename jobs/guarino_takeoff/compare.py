import openpyxl, pickle
p=r"C:\Users\jason\Downloads\Guarino Residence V3 - Estimate Items.xlsx"
wb=openpyxl.load_workbook(p,data_only=True); ws=wb["Estimate"]
LEAF={"MATERIAL","LABOR","SUBCONTRACTOR","ALLOWANCE","FEE","EQUIPMENT"}
# distribution of markup % he used
from collections import Counter
mks=Counter(); raw=0; amt=0
for r in range(2,ws.max_row+1):
    if str(ws.cell(r,3).value) in LEAF:
        q=ws.cell(r,7).value or 0; h=ws.cell(r,8).value or 0; mk=ws.cell(r,11).value; n=ws.cell(r,14).value or 0
        if n:
            mks[mk]+=1; raw+=q*h; amt+=n
print("his markup% distribution:",dict(mks))
print(f"his raw (G*H)={raw:,.0f}  his Amount={amt:,.0f}  blended markup={amt/raw-1:+.1%}")
myroll,mytot=pickle.load(open("results.pkl","rb"))[1], None
NOTES,myroll,R=pickle.load(open("results.pkl","rb"))
jroll,jtot,jrows=pickle.load(open("jason.pkl","rb"))
# unify keys
keys=sorted(set(myroll)|set(jroll), key=lambda k:-(jroll.get(k,0)))
print(f"\n{'CATEGORY':<42}{'MINE':>11}{'JASON':>11}{'Δ$':>11}{'Δ%':>8}")
for k in keys:
    m=myroll.get(k,0); j=jroll.get(k,0); d=m-j
    pc=(d/j*100) if j else 0
    print(f"{k:<42}{m:>11,.0f}{j:>11,.0f}{d:>11,.0f}{pc:>7.0f}%")
print(f"{'-'*42}{'-'*11}{'-'*11}")
print(f"{'TOTAL (marked-up line subtotal)':<42}{R['sub']:>11,.0f}{jtot:>11,.0f}{R['sub']-jtot:>11,.0f}{(R['sub']/jtot-1)*100:>7.1f}%")
print(f"his raw cost: {raw:,.0f}   my raw cost: {R['raw']:,.0f}   Δ {(R['raw']/raw-1)*100:+.1f}%")
