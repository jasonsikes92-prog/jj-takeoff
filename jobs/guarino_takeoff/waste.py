import openpyxl
p=r"C:\Users\jason\.claude\skills\jnj-estimate-takeoff\templates\measurements-template.xlsx"
wb=openpyxl.load_workbook(p,data_only=True)
print("SHEETS:",wb.sheetnames)
for sh in wb.sheetnames:
    ws=wb[sh]
    print(f"\n##### {sh}  ({ws.dimensions})")
    for r in range(1,min(ws.max_row,60)+1):
        vals=[ws.cell(r,c).value for c in range(1,min(ws.max_column,8)+1)]
        if any(v not in (None,"") for v in vals):
            print(r,"|"," | ".join("" if v is None else str(v)[:22] for v in vals))
