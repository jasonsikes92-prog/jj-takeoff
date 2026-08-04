import openpyxl
g=openpyxl.load_workbook(r"C:\Users\jason\OneDrive\Desktop\Claude\guarino_takeoff\J & J Pre-Estimate Intake Form.xlsx")["Guarino (example)"]
for r in range(5,g.max_row+1):
    a=g.cell(r,1).value; b=g.cell(r,2).value
    if a and not b: print(f"[{a}]")            # section header
    elif a: print(f"   {a[:40]:<40} = {b}")
