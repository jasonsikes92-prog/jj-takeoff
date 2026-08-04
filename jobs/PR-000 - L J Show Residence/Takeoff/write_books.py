# -*- coding: utf-8 -*-
import sys, json, math, os, collections
sys.path.insert(0, r'C:\Users\jason\.claude\skills\jnj-estimate-takeoff\tools')
import jnj_takeoff as J
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

RB = json.load(open(r'C:\Users\jason\.claude\skills\jnj-estimate-takeoff\reference\rate_book.json'))
LINES = RB['lines']
matched = json.load(open('matched.json'))
MARKUP = {"MATERIAL": 15.0, "LABOR": 7.0, "SUBCONTRACTOR": 7.0,
          "EQUIPMENT": 7.0, "FEE": 15.0, "ALLOWANCE": 8.0}

# client-facing descriptions (plain English, no takeoff math, no internal tags, no bare qty echo)
DESC = {
 'Shingle Roof - Turnkey': 'Architectural asphalt shingle roof on the main 6:12 roof planes, turnkey including underlayment, valley ice-and-water, ridge vent and starter.',
 'Metal Roofing - Standing Seam': 'Standing-seam metal roofing on the low-slope porch, entry and shed roofs, where the pitch is below the minimum for asphalt shingles.',
 'Siding (Vertical)': 'Fiber-cement board-and-batten siding, installed and ready for paint.',
 'Siding (Fiber Cement, Horizontal)': 'Fiber-cement lap siding, installed and ready for paint.',
 'Stone Material': 'Natural stone veneer to the exterior wainscot, entry mass, piers and chimney. See selections for the stone.',
 'Stone Labor': 'Mason labor to set the exterior stone veneer, including mortar, ties and flashing.',
 'Porch Ceilings': 'Tongue-and-groove wood ceiling to the covered front and rear porches, finished ready for stain or paint.',
 'Fascia (Metal and Vinyl)': 'Fascia and soffit to the full roof edge, including all eaves, rakes and porch returns.',
 'Corner Boards': 'Exterior corner boards at every outside corner, full wall height.',
 'Insulation - Spray Foam Roof Deck': 'Closed-cell spray foam insulation applied at the roof deck, creating a sealed conditioned attic.',
 'Insulation - Spray Foam Walls': 'Spray foam insulation to the exterior walls for a sealed, high-performance envelope.',
 'Drywall - Level 4': 'Hung, taped and finished drywall to a Level 4 finish throughout the house and garage, ready for paint.',
 'Labor - Monolithic slab': 'Place and finish the monolithic slab, including the thickened perimeter edge.',
 'Framing Lumber': 'Complete framing lumber package for the house, garage and covered porches.',
 'Framing Labor - SF Price': 'Framing labor for walls, roof structure and all covered areas.',
 'HVAC - Air Handler and Condensing Units': 'One high-efficiency heating and cooling system sized for the sealed spray-foam envelope, including air handler and condensing unit.',
 'Electrical - SF Price': 'Complete electrical rough-in and trim for the house and garage, including devices, panel and fixtures wiring.',
 'Well - 300 FT depth': 'Drilled water well to approximately 300 feet, including casing, pump and pressure tank.',
 'Sewer: Septic System, Plan, & Permit': 'Septic system including the site evaluation, permit, tank, field lines and installation.',
 'Trim Labor': 'Interior finish carpentry labor including door and window casing, base, and closet systems.',
 'Base molding - Choose size': 'Painted base molding throughout the heated areas. See selections for the profile.',
 'Granite or Quartz  - Kitchen': 'Level 3 granite or quartz kitchen countertops, including the island, with polished edges and sink cutout.',
 'Appliance Package  - Kitchen': 'Kitchen appliance package allowance. See selections for the final appliance schedule.',
 'Fireplace - Allowance': 'Great-room fireplace allowance, including firebox, flue and gas connection.',
 'Freize at roof': 'Frieze board at the roof line around the house, installed and ready for paint.',
 'Garage Wrap': 'Trimmed and wrapped surround to each garage door opening, including the header.',
 'Wrap Porch Beams': 'Wrapped and trimmed porch beams to the covered front and rear porches.',
 'Count Each Fixture Opening': 'Plumbing rough-in and fixture set for each fixture location throughout the house.',
 'Blower Door / HERS Test': 'Blower-door and energy-code performance testing of the finished envelope.',
 'Mudroom Built-ins / Drop Zone - Allowance': 'Mudroom built-ins including bench, cubbies and lockers. See selections for the finish.',
 'Punch-Out Labor - Allowance': 'End-of-job punch list labor to complete final adjustments and touch-ups before handover.',
}
GENERIC = {
 'ALLOWANCE': 'Allowance. Final cost is set by the selection made; any difference is reconciled by change order.',
}

def desc_for(l):
    n = str(l.get('name', ''))
    if n in DESC:
        return DESC[n]
    d = str(l.get('description') or '').strip()
    if d:
        return d
    ct = str(l.get('cost_type', ''))
    if ct == 'ALLOWANCE':
        return '%s. %s' % (n, GENERIC['ALLOWANCE'])
    return '%s, supplied and installed.' % n

rows = []
for i, qty, note in matched:
    l = LINES[i]
    ct = str(l.get('cost_type', 'MATERIAL')).upper()
    uc = float(l.get('unit_cost') or 0)
    if ct == 'ASSEMBLY':
        qty = 0.0
    pct = float(l.get('markup_pct') or MARKUP.get(ct, 15.0))
    builder = qty * uc
    mk = builder * pct / 100.0
    rows.append({
        'Name': l.get('name'), 'Cost type': ct, 'Cost code': l.get('cost_code'),
        'Cost title': l.get('cost_title'), 'Quantity': round(qty, 2),
        'Unit': l.get('unit'), 'Unit cost': round(uc, 4), 'Markup': round(mk, 2),
        'Group': l.get('group'), 'Description': desc_for(l),
        '_builder': round(builder, 2), '_amount': round(builder + mk, 2),
        '_pct': pct, '_note': note,
    })

direct = sum(r['_builder'] for r in rows)
markup = sum(r['Markup'] for r in rows)
subtotal = direct + markup
HEATED, UNDER = 3245.7, 5529.8
print('lines: %d' % len(rows))
print('Builder cost + allowances (direct): $%s' % format(round(direct), ','))
print('Per-line markup:                    $%s' % format(round(markup), ','))
print('ESTIMATE SUBTOTAL:                  $%s' % format(round(subtotal), ','))
print('   $/under-roof SF  %.2f   (%.0f SF)' % (subtotal / UNDER, UNDER))
print('   $/heated SF      %.2f   (%.0f SF)' % (subtotal / HEATED, HEATED))

by = collections.defaultdict(float)
for r in rows:
    by[str(r['Cost title'] or r['Group'])] += r['_amount']
print('\nCATEGORY ROLLUP (incl. waste + per-line markup):')
for k, v in sorted(by.items(), key=lambda t: -t[1]):
    print('   %-42s $%10s   %5.1f%%' % (k[:42], format(round(v), ','), 100 * v / subtotal))

# ---------- Buildern import ----------
COLS = ['Name', 'Cost type', 'Cost code', 'Cost title', 'Quantity', 'Unit', 'Unit cost', 'Markup', 'Group', 'Description']
wb = Workbook(); ws = wb.active; ws.title = 'Import'
ws.append(COLS)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='5C90AC')
for r in rows:
    ws.append([r[c] for c in COLS])
for col, w in zip('ABCDEFGHIJ', [44, 15, 10, 26, 11, 10, 11, 11, 34, 96]):
    ws.column_dimensions[col].width = w
OUT = r'C:\Users\jason\OneDrive\Desktop\PR-000 - L J Show Residence'
os.makedirs(OUT, exist_ok=True)
os.makedirs(os.path.join(OUT, 'Estimate'), exist_ok=True)
p1 = os.path.join(OUT, 'Estimate', 'L J Show Residence - Buildern Import.xlsx')
wb.save(p1)

# ---------- working workbook ----------
wb2 = Workbook()
e = wb2.active; e.title = 'Estimate'
e.append(['Cost code', 'Group', 'Name', 'Cost type', 'Quantity', 'Unit', 'Unit cost',
          'Builder cost', 'Markup %', 'Markup $', 'Amount', 'Description'])
for c in e[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='5C90AC')
for r in rows:
    e.append([r['Cost code'], r['Group'], r['Name'], r['Cost type'], r['Quantity'], r['Unit'],
              r['Unit cost'], r['_builder'], r['_pct'], r['Markup'], r['_amount'], r['Description']])
e.append([]); e.append(['', '', 'BUILDER COST + ALLOWANCES (DIRECT)', '', '', '', '', round(direct, 2), '', round(markup, 2), round(subtotal, 2)])
e.append(['', '', 'NOTE: Overhead/Profit and Contingency are applied in the Buildern SUMMARY, not as estimate lines.'])
for col, w in zip('ABCDEFGHIJKL', [10, 32, 42, 15, 11, 9, 11, 14, 9, 12, 13, 90]):
    e.column_dimensions[col].width = w

m = wb2.create_sheet('Measurements')
m.append(['Cost code', 'Item', 'Calc / basis', 'Type', 'Waste %', 'Qty', 'Unit', 'Plan sheet'])
for c in m[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='5C90AC')
MEAS = [
 ['', 'Heated area', 'Vector area-polygon shoelace at 9.0 pts/ft, verified by raster colour-mask and 3 independent footprint traces', 'MEASURED', 0, 3245.7, 'SF', 'S2 SQFT / S5'],
 ['', 'Garage', 'Same method; polygon reproduces the drawn area to 0.05%', 'MEASURED', 0, 1038.5, 'SF', 'S2 SQFT'],
 ['', 'Rear covered porch', 'Same method; roofed by the 1:12 shed planes on S6', 'MEASURED', 0, 1168.8, 'SF', 'S2 / S6'],
 ['', 'Front covered porch', 'Same method; roofed by the 4:12 entry shed', 'MEASURED', 0, 76.8, 'SF', 'S2 / S6'],
 ['', 'TOTAL UNDER ROOF (framing basis)', 'heated + garage + both covered porches', 'DERIVED', 0, 5529.8, 'SF', '-'],
 ['4.15', 'Slab area', 'Under-roof footprint, monolithic slab confirmed by Jason', 'MEASURED', 0, 5529.8, 'SF', 'S4 SLAB PLAN'],
 ['4.15', 'Thickened slab edge', 'Exterior envelope perimeter traced off the area polygons', 'MEASURED', 0, 387.69, 'LF', 'S4 / S2'],
 ['6', 'Roof footprint', 'Vector roof-outline trace at verified 18.0 pts/ft', 'MEASURED', 0, 6159.0, 'SF', 'S6 ROOF PLAN'],
 ['6', 'Roof overhang', 'Roof-edge-to-wall distance histogram (mode 18in) and cross-checked by footprint difference (1.57 ft)', 'MEASURED', 0, 1.5, 'FT', 'S6'],
 ['6', 'Shingle roof surface 6:12', 'Face decomposition by pitch zone; 4,346.8 SF footprint x 1.1180 pitch factor', 'MEASURED', 15, 5588.7, 'SF', 'S6'],
 ['6', 'Standing-seam metal surface', '1:12 / 2:12 / 3:12 / 4:12 zones; below shingle minimum pitch', 'MEASURED', 15, 2095.9, 'SF', 'S6 / S8 / S9'],
 ['5.1', 'Board-and-batten siding', 'Per-wall and per-gable measurement; each of 32 walls sampled on the elevation that sees it', 'MEASURED', 10, 3649.8, 'SF', 'S8 / S9 / S5'],
 ['5.1', 'Lap siding', 'Same per-wall method', 'MEASURED', 10, 1823.8, 'SF', 'S8 / S9 / S5'],
 ['17.2', 'Stone veneer', 'Wainscot band plus entry mass and chimney; measured per elevation face', 'MEASURED', 10, 618.2, 'SF', 'S8 / S9'],
 ['5.1', 'Fascia and soffit', 'Roof-plan plan perimeter 402.0 LF plus 54.2 LF rake slope uplift across 13 measured gables', 'MEASURED', 0, 456.2, 'LF', 'S6 + S8/S9'],
 ['5.1', 'Corner boards', '32 envelope direction changes x 11.05 ft wall height', 'MEASURED', 0, 354.0, 'LF', 'S2 / S8'],
 ['5.1', 'Covered porch T&G ceiling', 'Both covered porch areas', 'MEASURED', 10, 1370.2, 'SF', 'S2 / S6'],
 ['10.05', 'Drywall net surface', 'Room-by-room: 41 measured rooms, perimeter x that room ceiling height, plus ceiling area', 'MEASURED', 10, 22181.6, 'SF', 'S5'],
 ['10', 'Spray foam roof deck', 'True roof surface area by pitch zone', 'MEASURED', 10, 7350.4, 'SF', 'S6'],
 ['10', 'Spray foam walls', 'Exterior wall band area', 'MEASURED', 10, 4721.2, 'SF', 'S2 / S8 / S9'],
 ['13', 'Windows', 'Counted off the floor plan tags and reconciled against all four elevations', 'COUNTED', 0, 22, 'EA', 'S5 / S8 / S9'],
 ['13.3', 'Interior doors', 'Counted off floor plan door tags', 'COUNTED', 0, 27, 'EA', 'S5'],
 ['8', 'Recessed cans', 'Counted off the electrical plan symbols', 'COUNTED', 0, 59, 'EA', 'S7'],
 ['9', 'HVAC systems', '3,245.7 heated / 800 = 4.06 tons, under the 5-ton per-unit maximum', 'DERIVED', 0, 1, 'EA', '-'],
]
for r in MEAS:
    m.append(r)
for col, w in zip('ABCDEFGH', [10, 34, 96, 12, 9, 12, 8, 20]):
    m.column_dimensions[col].width = w

a = wb2.create_sheet('Allowances')
a.append(['Room / scope', 'Line', 'Allowance $', 'Note'])
for c in a[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='5C90AC')
for r in rows:
    if r['Cost type'] == 'ALLOWANCE':
        a.append([r['Group'], r['Name'], r['_builder'], r['Description'][:120]])
for col, w in zip('ABCD', [34, 46, 14, 90]):
    a.column_dimensions[col].width = w

s = wb2.create_sheet('Selections')
s.append(['Scope', 'Shown on plans', 'Selection needed', 'Status'])
for c in s[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='5C90AC')
SEL = [
 ['Roof - main planes', 'Architectural shingle hatch on all 6:12 planes', 'Shingle brand / colour', 'CONFIRM'],
 ['Roof - low slope', 'Standing-seam ribs on 1:12/2:12/3:12/4:12 planes', 'Metal profile, gauge and colour; separate vendor', 'CONFIRM'],
 ['Siding', 'Board-and-batten dominant, lap on front gables', 'Hardie product and colour; lap priced at the updated rate', 'CONFIRMED w/ Jason'],
 ['Stone', 'Wainscot, entry mass, chimney, porch piers', 'Stone type and colour', 'CONFIRM'],
 ['Windows', '22 units, mostly fixed and single-hung', 'Brand, colour, grid pattern', 'CONFIRM'],
 ['Kitchen counters', 'Island plus perimeter runs on E5-E8', 'Level 3 granite or quartz', 'CONFIRM'],
 ['Flooring', 'Wood-look plank throughout living, tile in baths', 'Product per room', 'CONFIRM'],
 ['Tubs', 'Freestanding tub shown in the master bath render', 'Tub model per bath', 'CONFIRM'],
 ['Appliances', 'Range, wall oven, hood, fridge on E5-E8', 'Appliance tier', 'CONFIRM'],
 ['Foundation', 'Sheet 4 SLAB PLAN', 'Monolithic slab', 'CONFIRMED w/ Jason'],
 ['Water / sewer', 'Rural acreage, Spalding County', 'Well + septic', 'CONFIRMED w/ Jason'],
 ['Insulation', 'Sealed envelope', 'Spray foam roof deck + spray foam walls', 'CONFIRMED w/ Jason'],
]
for r in SEL:
    s.append(r)
for col, w in zip('ABCD', [26, 52, 52, 22]):
    s.column_dimensions[col].width = w

t = wb2.create_sheet('Assumptions to Confirm')
t.append(['#', 'Item', 'What I assumed', 'Why it matters', 'Impact'])
for c in t[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='C00000')
ASSUME = [
 [1, 'NO SITE PLAN in the set', 'Driveway 300 LF x 12 ft, silt fence 900 LF, permanent power 300 LF, 4 days clearing', 'These are pure placeholders - the set has no site plan and the plat is a 2023 lot-division survey, not a grading plan', 'HIGH - can move +/- $30k'],
 [2, 'Well distance from house', '150 LF of well electrical', 'Priced per LF; needs the real distance', 'MED'],
 [3, 'Dumpster count', '6 pulls', 'Judgement call based on how wooded the lot is', 'LOW'],
 [4, 'Build duration', '10 months (J&J default)', 'Drives temp toilets and temp utilities', 'LOW'],
 [5, 'Porch column count', '10 columns wrapped', 'Counted approximately off the rear elevation; each 12x12 cedar wrap is $480', 'MED'],
 [6, 'Drywall quantity', '22,182 SF ordered on a 20,165 SF net measure', 'Room-by-room across 41 measured rooms at 10-12 ft ceilings. Ratio is above a typical house because of the tall plates and the cut-up plan - worth a sanity check against a Burns-style quote', 'HIGH'],
 [7, 'Interior door count', '27 doors at 8 ft', 'Counted off plan tags; the engine door counter expects 6-8 tags so I counted the 8-0 tags manually', 'MED'],
 [8, 'Cabinet run LF', 'Kitchen 24 lower / 18 upper / 12 island / 8 tall', 'Read off the E5-E8 kitchen elevations, not a full room-by-room cabinet takeoff', 'MED'],
 [9, 'Bath fixture count', '5 full baths + 1 powder', 'From the room labels on sheet 5', 'MED'],
 [10, 'Stone vs siding split', 'Stone 562 SF', 'Measured per elevation face. Stone is mason scope and excluded from the siding sub', 'MED'],
 [11, 'Gable count', '13 gables, 1,247 SF', 'Rake-paired triangles plus unpaired rakes carried rather than dropped, with the 45-degree wing corrected', 'MED'],
 [12, 'Pre-con soft costs', 'Plan design, engineering, boundary survey and erosion plan all carried', 'Some may already be done - confirm which to drop', 'MED'],
]
for r in ASSUME:
    t.append(r)
for col, w in zip('ABCDE', [5, 30, 60, 78, 18]):
    t.column_dimensions[col].width = w
for sh in wb2:
    sh.freeze_panes = 'A2'
p2 = os.path.join(OUT, 'Estimate', 'L J Show Residence - Estimate.xlsx')
wb2.save(p2)
print('\nWROTE:\n  %s\n  %s' % (p1, p2))
json.dump({'direct': direct, 'markup': markup, 'subtotal': subtotal,
           'per_under_roof_sf': subtotal / UNDER, 'per_heated_sf': subtotal / HEATED,
           'rollup': {k: round(v, 2) for k, v in by.items()}}, open('totals.json', 'w'), indent=1)
